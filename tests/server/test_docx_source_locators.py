import io
import zipfile

import pytest

from server.services import extract, ingest, input_formats


def document(body):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", body)
        archive.writestr("word/_rels/document.xml.rels", '<Relationships><Relationship Target="http://127.0.0.1/private"/></Relationships>')
    return output.getvalue()


BODY = '''<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>
<w:p><w:r><w:t>First</w:t><w:tab/><w:t>line</w:t><w:br/><w:t>next</w:t></w:r></w:p>
<w:p/>
<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Table cell</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
<w:p><w:r><w:instrText>DO NOT EXECUTE FIELD</w:instrText><w:t>Cached display</w:t></w:r></w:p>
</w:body></w:document>'''


def test_word_tables_and_empty_paragraphs_preserve_source_positions():
    text, truncated = input_formats.read_structured("REPORT.DOCX", document(BODY))
    assert text == ("[word/document.xml#paragraph=1] First\tline\nnext\n"
                    "[word/document.xml#paragraph=3] Table cell\n"
                    "[word/document.xml#paragraph=4] Cached display")
    assert not truncated
    assert ingest._extract_file("report.docx", document(BODY)) == text


@pytest.mark.parametrize("body", [b"<invalid>", b'<!DOCTYPE x [<!ENTITY a "secret">]><x>&a;</x>', BODY.encode("utf-16"), b"<document/>"])
def test_word_invalid_xml_and_entities_fail_closed(body):
    with pytest.raises(input_formats.InputError, match="inputs.invalid"):
        input_formats.read_structured("report.docx", document(body))


def test_word_output_limit_is_explicit(monkeypatch):
    monkeypatch.setattr(input_formats, "MAX_TEXT", 60)
    text, truncated = input_formats.read_structured("report.docx", document(BODY))
    assert truncated and "paragraph=1" in text and "Table cell" not in text
    assert '"extraction_truncated": true' in ingest._extract_file("report.docx", document(BODY))


async def test_ephemeral_word_keeps_locators_and_does_not_compress(monkeypatch):
    async def forbidden(*args):
        raise AssertionError("Source-addressed Word must not invoke compression")
    monkeypatch.setattr(ingest, "_compress", forbidden)
    text, truncated = await extract.extract_text(filename="report.docx", data=document(BODY), compress=True)
    assert "[word/document.xml#paragraph=3] Table cell" in text
    assert not truncated


def test_nested_textbox_paragraph_is_not_duplicated():
    body = BODY.replace("<w:t>First</w:t>", "<w:t>First</w:t><w:txbxContent><w:p><w:r><w:t>Textbox</w:t></w:r></w:p></w:txbxContent>")
    text, _ = input_formats.read_structured("report.docx", document(body))
    assert text.count("Textbox") == 1
    assert "[word/document.xml#paragraph=2] Textbox" in text


def test_real_word_package_reads_body_and_table():
    from docx import Document
    doc = Document()
    doc.add_paragraph("正文 body")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "表格 table"
    output = io.BytesIO()
    doc.save(output)
    text, truncated = input_formats.read_structured("actual.docx", output.getvalue())
    assert text == "[word/document.xml#paragraph=1] 正文 body\n[word/document.xml#paragraph=2] 表格 table"
    assert not truncated


@pytest.mark.parametrize("removed_tag", ["del", "moveFrom"])
def test_revision_source_text_is_not_presented_as_current_body(removed_tag):
    body = f'''<w:document xmlns:w="{input_formats.NS['w']}"><w:body>
    <w:{removed_tag}><w:p><w:r><w:t>Old deadline Friday</w:t></w:r></w:p></w:{removed_tag}>
    <w:p><w:{removed_tag}><w:r><w:t>Old inline deadline</w:t></w:r></w:{removed_tag}>
    <w:ins><w:r><w:t>New deadline Monday</w:t></w:r></w:ins></w:p>
    <w:moveTo><w:p><w:r><w:t>Current moved paragraph</w:t></w:r></w:p></w:moveTo>
    </w:body></w:document>'''
    text, truncated = input_formats.read_structured("revision.docx", document(body))
    assert "Old deadline" not in text and "Old inline" not in text
    assert "[word/document.xml#paragraph=2] New deadline Monday" in text
    assert "[word/document.xml#paragraph=3] Current moved paragraph" in text
    assert "tracked revisions" in text
    assert "not a revision-history comparison" in text
    assert not truncated


async def test_revision_policy_matches_ephemeral_and_persistent_ingest():
    body = BODY.replace("<w:t>First</w:t>", '<w:del><w:r><w:delText>Removed secret</w:delText></w:r></w:del>'
                        '<w:moveFrom><w:txbxContent><w:p><w:r><w:t>Old textbox</w:t></w:r></w:p></w:txbxContent></w:moveFrom>'
                        '<w:ins><w:r><w:t>Current</w:t></w:r></w:ins>')
    data = document(body)
    direct, _ = input_formats.read_structured("revision.docx", data)
    ephemeral, truncated = await extract.extract_text(filename="revision.docx", data=data)
    assert ephemeral == direct == ingest._extract_file("revision.docx", data)
    assert "Removed secret" not in direct and "Old textbox" not in direct
    assert "paragraph=4] Table cell" in direct
    assert "Current" in direct and not truncated


def test_revision_disclosure_cannot_be_silently_dropped_at_text_limit(monkeypatch):
    body = BODY.replace("<w:t>First</w:t>", "<w:ins><w:r><w:t>New</w:t></w:r></w:ins>")
    monkeypatch.setattr(input_formats, "MAX_TEXT", 60)
    text, truncated = input_formats.read_structured("revision.docx", document(body))
    assert text == "" and truncated
