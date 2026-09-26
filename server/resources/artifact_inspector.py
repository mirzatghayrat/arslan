"""Fixed, isolated parser entry point. No user code or network access is needed."""
import csv
import io
import json
from pathlib import Path
import sys
import warnings
import zipfile
from xml.etree import ElementTree

MAX_EXPANDED = 50 * 1024 * 1024


def inspect(data, suffix):
    if not data:
        return {"status": "failed", "code": "artifact_empty"}
    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tiff"}:
        from PIL import Image
        Image.MAX_IMAGE_PIXELS = 25_000_000
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(io.BytesIO(data)) as image:
                size = {"width": image.width, "height": image.height}
                image.verify()
        return {"status": "passed", "code": "image_parsed", **size}
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data), strict=True)
        if reader.is_encrypted:
            return {"status": "not_run", "code": "encrypted_artifact"}
        if not 0 < len(reader.pages) <= 2000:
            return {"status": "failed", "code": "artifact_page_limit"}
        for page in reader.pages:
            _ = page.mediabox
            contents = page.get_contents()
            if contents is not None and len(contents.get_data()) > MAX_EXPANDED:
                return {"status": "failed", "code": "artifact_expansion_limit"}
        return {"status": "passed", "code": "pdf_parsed", "pages": len(reader.pages)}
    if suffix in {".docx", ".pptx", ".xlsx"}:
        required = {".docx": "word/document.xml", ".pptx": "ppt/presentation.xml", ".xlsx": "xl/workbook.xml"}[suffix]
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            items = archive.infolist()
            if len(items) > 2000 or sum(item.file_size for item in items) > MAX_EXPANDED:
                return {"status": "failed", "code": "artifact_expansion_limit"}
            if required not in archive.namelist() or "[Content_Types].xml" not in archive.namelist():
                return {"status": "failed", "code": "office_container_invalid"}
            for item in items:
                content = archive.read(item)
                if item.filename.endswith((".xml", ".rels")):
                    if b"<!DOCTYPE" in content.upper() or b"<!ENTITY" in content.upper():
                        return {"status": "failed", "code": "xml_entities_denied"}
                    ElementTree.fromstring(content)
        if suffix == ".docx":
            from docx import Document
            document = Document(io.BytesIO(data))
            _ = document.paragraphs, document.tables, document.sections
        elif suffix == ".pptx":
            from pptx import Presentation
            presentation = Presentation(io.BytesIO(data))
            for slide in presentation.slides:
                _ = slide.shapes
        else:
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            try:
                for sheet in workbook:
                    if sheet.max_row and sheet.max_column and sheet.max_row * sheet.max_column > 1_000_000:
                        return {"status": "failed", "code": "artifact_cell_limit"}
                    for _ in sheet.iter_rows():
                        pass
            finally:
                workbook.close()
        return {"status": "passed", "code": "office_document_parsed"}
    if suffix in {".txt", ".md", ".json", ".csv", ".tsv", ".html", ".css", ".js", ".ts", ".py", ".svg"}:
        text = data.decode("utf-8-sig")
        if suffix == ".json":
            json.loads(text)
        elif suffix in {".csv", ".tsv"}:
            for row in csv.reader(io.StringIO(text), delimiter="\t" if suffix == ".tsv" else ",", strict=True):
                if len(row) > 10_000:
                    return {"status": "failed", "code": "artifact_column_limit"}
        elif suffix == ".svg":
            if "<!DOCTYPE" in text.upper() or "<!ENTITY" in text.upper():
                return {"status": "failed", "code": "xml_entities_denied"}
            ElementTree.fromstring(text)
        return {"status": "passed", "code": "text_parsed", "characters": len(text)}
    return {"status": "not_run", "code": "artifact_parser_unavailable"}


if __name__ == "__main__":
    try:
        answer = inspect(Path("input.bin").read_bytes(), sys.argv[1].lower())
    except ImportError:
        answer = {"status": "not_run", "code": "artifact_parser_unavailable"}
    except Exception:
        answer = {"status": "failed", "code": "artifact_parse_failed"}
    print(json.dumps(answer))
